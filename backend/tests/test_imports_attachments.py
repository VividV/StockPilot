import io
from datetime import date, timedelta
from decimal import Decimal

from openpyxl import Workbook, load_workbook


def xlsx(rows: list[list]) -> bytes:
    wb = Workbook()
    for r in rows:
        wb.active.append(r)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def upload(api, kind, name, data):
    return api.post(f"/api/imports/{kind}/preview", files={"file": (name, data)})


def test_templates_download_and_round_trip(org):
    r = org.owner.get("/api/imports/products/template?format=xlsx")
    assert r.status_code == 200 and r.content.startswith(b"PK")
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Products", "How to fill this in"]
    assert wb.active["A1"].value == "SKU *"
    # the untouched template (header + example row) has no data rows
    r2 = upload(org.owner, "products", "t.xlsx", r.content)
    assert r2.status_code == 422 and "No data rows" in r2.json()["detail"]
    csv_t = org.owner.get("/api/imports/customers/template?format=csv")
    assert csv_t.content.decode("utf-8-sig").splitlines()[0].startswith("Name,Contact person")


def test_product_import_preview_errors_fix_and_commit(org):
    sup = org.owner.ok("post", "/api/suppliers", json={"name": "Acme Textiles"})
    existing = org.product(sku="OLD-1", name="Old name", barcode="111")
    data = xlsx([
        ["SKU", "Name", "Barcode", "Category", "Unit", "Cost price", "Selling price", "Reorder level", "Preferred supplier", "Track expiry"],
        ["NEW-1", "Scarf", "222", "Accessories", "pcs", "1500", "4000", "5", "Acme Textiles", "no"],
        ["OLD-1", "Old name, renamed", "", "", "", "", "9.99", "", "", ""],
        ["NEW-2", "", "111", "", "litres??", "abc", "", "", "Nobody Ltd", "maybe"],
        ["NEW-1", "Duplicate", "", "", "", "", "", "", "", ""],
    ])
    r = upload(org.owner, "products", "products.xlsx", data)
    assert r.status_code == 200, r.text
    rep = r.json()
    assert rep["summary"] == {"total": 4, "create": 3, "update": 1, "with_errors": 2}
    rows = rep["rows"]
    assert rows[0]["row"] == 2 and rows[0]["errors"] == [] and "New category 'Accessories'" in rows[0]["notes"][0]
    assert rows[1]["action"] == "update"
    bad = {e["field"] for e in rows[2]["errors"]}
    assert bad == {"name", "barcode", "unit", "cost_price", "supplier", "track_expiry"}
    assert "Already used by product OLD-1" in [e["message"] for e in rows[2]["errors"] if e["field"] == "barcode"][0]
    assert rows[3]["errors"][0]["message"] == "Same SKU as row 2"

    # committing with errors saves nothing
    r = org.owner.post("/api/imports/products/commit", json={"rows": [x["data"] for x in rows]})
    assert r.status_code == 422 and "Nothing was imported" in r.json()["detail"]
    assert org.owner.ok("get", "/api/products?q=Scarf")["total"] == 0

    # fix on screen: correct row 4, drop the duplicate
    fixed = [rows[0]["data"], rows[1]["data"],
             {**rows[2]["data"], "name": "Gloves", "barcode": "333", "unit": "pcs", "cost_price": "800", "supplier": "", "track_expiry": "no"}]
    v = org.owner.ok("post", "/api/imports/products/validate", json={"rows": fixed})
    assert v["summary"]["with_errors"] == 0
    r = org.owner.ok("post", "/api/imports/products/commit", json={"rows": fixed})
    assert r["created"] == 2 and r["updated"] == 1
    scarf = org.owner.ok("get", "/api/products?q=Scarf")["items"][0]
    assert scarf["category"]["name"] == "Accessories" and scarf["preferred_supplier"]["id"] == sup["id"]
    assert Decimal(scarf["selling_price"]) == 4000 and Decimal(scarf["reorder_level"]) == 5
    old = org.owner.ok("get", f"/api/products/{existing['id']}")["product"]
    assert old["name"] == "Old name, renamed" and Decimal(old["selling_price"]) == Decimal("9.99") and old["barcode"] == "111"
    assert org.owner.ok("get", "/api/audit-logs?action=import.products")["total"] == 1


def test_csv_with_semicolons_and_partner_upsert(org):
    org.owner.ok("post", "/api/customers", json={"name": "Lekki Mart", "phone": "1"})
    csv_bytes = "Name;Contact person;Email;Phone\nLekki Mart;Femi;femi@lekki.example;2\nNew Shop;;bad-email;\n".encode("utf-8")
    rep = upload(org.owner, "customers", "c.csv", csv_bytes).json()
    assert [r["action"] for r in rep["rows"]] == ["update", "create"]
    assert rep["rows"][1]["errors"] == [{"field": "email", "message": "Not a valid email address"}]
    rows = [rep["rows"][0]["data"], {**rep["rows"][1]["data"], "email": ""}]
    r = org.owner.ok("post", "/api/imports/customers/commit", json={"rows": rows})
    assert (r["created"], r["updated"]) == (1, 1)
    lekki = org.owner.ok("get", "/api/customers?q=Lekki")["items"][0]
    assert lekki["phone"] == "2" and lekki["contact_name"] == "Femi"


def test_opening_stock_import(org):
    shop = org.warehouse("Ikeja Shop")
    plain = org.product(sku="P1", base_unit_id=org.units["kg"], cost_price=3)
    med = org.product(sku="M1", track_expiry=True)
    lap = org.product(sku="L1", track_serials=True)
    exp = (date.today() + timedelta(days=40)).isoformat()
    data = xlsx([
        ["SKU", "Location", "Quantity", "Cost per unit", "Batch number", "Expiry date", "Serial numbers"],
        ["P1", "Main Store", "12.5", "", "", "", ""],
        ["M1", "ikeja shop", "10", "2", "LOT-1", exp, ""],
        ["L1", "MAIN", "2", "500", "", "", "SN-1; SN-2"],
        ["M1", "Main Store", "5", "", "", "", ""],
        ["ZZZ", "Nowhere", "-1", "", "", "", ""],
    ])
    rep = upload(org.owner, "opening-stock", "o.xlsx", data).json()
    errs = [{e["field"] for e in r["errors"]} for r in rep["rows"]]
    assert errs[:3] == [set(), set(), set()]
    assert errs[3] == {"batch_number", "expires_on"}
    assert errs[4] == {"sku", "location", "quantity"}
    rows = [r["data"] for r in rep["rows"][:3]]
    r = org.owner.ok("post", "/api/imports/opening-stock/commit", json={"rows": rows})
    assert r["created"] == 3
    assert org.on_hand(plain["id"]) == 12.5 and org.on_hand(med["id"]) == 10 and org.on_hand(lap["id"]) == 2
    mv = org.owner.ok("get", f"/api/movements?product_id={med['id']}")["items"][0]
    assert mv["movement_type"] == "opening_balance" and mv["batch_number"] == "LOT-1" and mv["location_name"] == "Ikeja Shop"
    serials = {s["serial"] for s in org.owner.ok("get", f"/api/serials?product_id={lap['id']}")["items"]}
    assert serials == {"SN-1", "SN-2"}
    # the same serials again would be duplicates -> all-or-nothing: nothing changes
    r = org.owner.post("/api/imports/opening-stock/commit", json={"rows": [rows[0], rows[2]]})
    assert r.status_code == 400 and org.on_hand(plain["id"]) == 12.5


def test_import_permissions_and_bad_files(org):
    viewer = org.user("Viewer")
    assert viewer.get("/api/imports/products/template").status_code == 403
    assert viewer.post("/api/imports/products/commit", json={"rows": [{"sku": "x", "name": "y"}]}).status_code == 403
    sales = org.user("Sales Staff")
    assert sales.post("/api/imports/customers/validate", json={"rows": [{"name": "x"}]}).status_code == 200
    assert sales.post("/api/imports/opening-stock/validate", json={"rows": [{"sku": "x"}]}).status_code == 403
    r = upload(org.owner, "products", "notes.pdf", b"%PDF-1.4")
    assert r.status_code == 422 and "Excel" in r.json()["detail"]
    r = upload(org.owner, "products", "x.xlsx", xlsx([["Name", "Price"], ["a", 1]]))
    assert r.status_code == 422 and "Missing column(s): SKU" in r.json()["detail"]


# ------------------------------------------------------------------ attachments
PNG = b"\x89PNG\r\n\x1a\n" + b"0" * 100


def test_attachments_upload_list_download_delete(org):
    p = org.product()
    r = org.owner.post("/api/attachments", data={"entity_type": "product", "entity_id": p["id"]},
                       files={"file": ("../../etc/photo.png", PNG, "image/png")})
    assert r.status_code == 201, r.text
    a = r.json()
    assert a["filename"] == "photo.png" and a["is_image"]
    lst = org.owner.ok("get", f"/api/attachments?entity_type=product&entity_id={p['id']}")
    assert [x["id"] for x in lst] == [a["id"]]
    f = org.owner.get(f"/api/attachments/{a['id']}/file")
    assert f.status_code == 200 and f.content == PNG and f.headers["content-type"] == "image/png"
    assert org.owner.ok("delete", f"/api/attachments/{a['id']}")["message"] == "File removed"
    assert org.owner.get(f"/api/attachments/{a['id']}/file").status_code == 404


def test_attachment_safety(org, make_org):
    p = org.product()
    form = {"entity_type": "product", "entity_id": p["id"]}
    r = org.owner.post("/api/attachments", data=form, files={"file": ("evil.html", b"<script>", "text/html")})
    assert r.status_code == 422
    r = org.owner.post("/api/attachments", data=form, files={"file": ("fake.png", b"<html>not an image", "image/png")})
    assert r.status_code == 422 and "real PNG" in r.json()["detail"]
    r = org.owner.post("/api/attachments", data=form, files={"file": ("big.pdf", b"%PDF" + b"0" * (10 * 1024 * 1024), "application/pdf")})
    assert r.status_code == 413
    ok = org.owner.post("/api/attachments", data=form, files={"file": ("note.pdf", b"%PDF-1.4 hello", "application/pdf")}).json()
    viewer = org.user("Viewer")
    assert viewer.get(f"/api/attachments/{ok['id']}/file").status_code == 200       # can see the product
    assert viewer.post("/api/attachments", data=form, files={"file": ("n.pdf", b"%PDF", "application/pdf")}).status_code == 403
    other = make_org()
    assert other.owner.get(f"/api/attachments/{ok['id']}/file").status_code == 404
    assert other.owner.post("/api/attachments", data=form, files={"file": ("n.pdf", b"%PDF", "application/pdf")}).status_code == 404
