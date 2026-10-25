"""Spreadsheet import for products, suppliers, customers and opening stock.

Flow: template -> upload -> validate (row-by-row errors) -> user fixes rows ->
validate again -> commit. Commit re-validates and applies everything in one
transaction: either every row is imported or none is.
"""
from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.models import (
    Category, Customer, Location, LocationKind, MovementType, Organization, Product, StockMovement, Supplier,
    UnitOfMeasure, Warehouse,
)
from app.services import inventory as inv
from app.services.industry import product_defaults

MAX_ROWS = 5000
EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")


@dataclass
class Col:
    key: str
    label: str
    required: bool = False
    help: str = ""
    example: str = ""


KINDS: dict[str, dict] = {
    "products": {
        "title": "Products", "match": "SKU",
        "columns": [
            Col("sku", "SKU", True, "Your unique item code. Existing SKUs are updated.", "TS-RED-M"),
            Col("name", "Name", True, "", "T-shirt, red, medium"),
            Col("barcode", "Barcode", False, "", "6151100000123"),
            Col("category", "Category", False, "Created if it doesn't exist", "Clothing"),
            Col("brand", "Brand", False, "", "Dye Lab"),
            Col("unit", "Unit", False, "pcs, box, ctn, kg, l, m… Defaults to pcs", "pcs"),
            Col("cost_price", "Cost price", False, "", "2500"),
            Col("selling_price", "Selling price", False, "", "6000"),
            Col("tax_rate", "Tax %", False, "", "7.5"),
            Col("reorder_level", "Reorder level", False, "Warn when stock falls to this", "10"),
            Col("reorder_qty", "Reorder quantity", False, "How many you usually order", "50"),
            Col("min_stock", "Minimum stock", False, "", ""),
            Col("supplier", "Preferred supplier", False, "Must match a supplier name", ""),
            Col("description", "Description", False, "", ""),
            Col("track_expiry", "Track expiry", False, "yes / no", "no"),
            Col("track_batches", "Track batches", False, "yes / no", "no"),
            Col("track_serials", "Track serial numbers", False, "yes / no", "no"),
            Col("is_active", "Active", False, "yes / no (default yes)", "yes"),
        ],
    },
    "suppliers": {
        "title": "Suppliers", "match": "name",
        "columns": [
            Col("name", "Name", True, "Existing suppliers with the same name are updated", "Acme Textiles Ltd"),
            Col("contact_name", "Contact person", False, "", "Ada Obi"),
            Col("email", "Email", False, "", "orders@acme.example"),
            Col("phone", "Phone", False, "", "+234 801 000 0000"),
            Col("address", "Address", False, "", "12 Allen Avenue, Ikeja"),
            Col("notes", "Notes", False, "", ""),
        ],
    },
    "customers": {
        "title": "Customers", "match": "name",
        "columns": [
            Col("name", "Name", True, "Existing customers with the same name are updated", "Lekki Mart"),
            Col("contact_name", "Contact person", False, "", "Femi Ade"),
            Col("email", "Email", False, "", "buying@lekkimart.example"),
            Col("phone", "Phone", False, "", "+234 804 000 0000"),
            Col("address", "Address", False, "", ""),
            Col("notes", "Notes", False, "", ""),
        ],
    },
    "opening-stock": {
        "title": "Opening stock", "match": None,
        "columns": [
            Col("sku", "SKU", True, "Must match a product", "TS-RED-M"),
            Col("location", "Location", True, "Warehouse / shop name or code", "Main Store"),
            Col("bin", "Shelf / bin", False, "Optional code of a shelf or bin inside it", ""),
            Col("quantity", "Quantity", True, "In the product's base unit", "24"),
            Col("unit_cost", "Cost per unit", False, "Defaults to the product's cost price", "2500"),
            Col("batch_number", "Batch number", False, "Required for batch-tracked products", ""),
            Col("expires_on", "Expiry date", False, "YYYY-MM-DD. Required when expiry is tracked", ""),
            Col("serials", "Serial numbers", False, "Separate with ; — one per unit", ""),
        ],
    },
}


def _norm(s) -> str:
    return re.sub(r"[^a-z0-9%]", "", str(s or "").lower())


# ------------------------------------------------------------------ templates
def template(kind: str, fmt: str) -> bytes:
    cols = KINDS[kind]["columns"]
    if fmt == "csv":
        buf = io.StringIO()
        w = csv.writer(buf)
        w.writerow([c.label for c in cols])
        w.writerow([c.example for c in cols])
        return ("﻿" + buf.getvalue()).encode("utf-8")
    wb = Workbook()
    ws = wb.active
    ws.title = KINDS[kind]["title"][:31]
    ws.append([c.label + (" *" if c.required else "") for c in cols])
    ws.append([c.example for c in cols])
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=1, column=i)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="F2C230" if c.required else "E8EEF6")
        ws.column_dimensions[get_column_letter(i)].width = max(14, len(c.label) + 4)
        ws.cell(row=2, column=i).font = Font(italic=True, color="6B7A90")
    ws.freeze_panes = "A2"
    help_ws = wb.create_sheet("How to fill this in")
    help_ws.append(["Column", "Required", "Notes"])
    for c in cols:
        help_ws.append([c.label, "Yes" if c.required else "", c.help])
    help_ws.append([])
    help_ws.append(["Row 2 is an example. Replace it with your own data or delete it."])
    for col, width in (("A", 24), ("B", 10), ("C", 60)):
        help_ws.column_dimensions[col].width = width
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# ------------------------------------------------------------------ parsing
class ImportFileError(Exception):
    pass


def parse_file(kind: str, filename: str, data: bytes) -> list[dict]:
    cols = KINDS[kind]["columns"]
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception:
            raise ImportFileError("This Excel file couldn't be opened. Save it as .xlsx and try again.")
        ws = wb.worksheets[0]
        raw = [[("" if v is None else v) for v in row] for row in ws.iter_rows(values_only=True)]
    elif name.endswith((".csv", ".txt")):
        text = None
        for enc in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                text = data.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        dialect = csv.excel
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        except csv.Error:
            pass
        raw = list(csv.reader(io.StringIO(text), dialect))
    else:
        raise ImportFileError("Upload an Excel (.xlsx) or CSV file. Download the template to start.")
    raw = [r for r in raw if any(str(v).strip() for v in r)]
    if not raw:
        raise ImportFileError("The file is empty")
    header = [_norm(h).rstrip("*") for h in raw[0]]
    lookup = {}
    for c in cols:
        for alias in (c.key, c.label):
            lookup[_norm(alias)] = c.key
    mapping = {i: lookup[h] for i, h in enumerate(header) if h in lookup}
    missing = [c.label for c in cols if c.required and c.key not in mapping.values()]
    if missing:
        raise ImportFileError(f"Missing column(s): {', '.join(missing)}. The first row must hold the column names from the template.")
    rows = []
    example = {c.key: c.example for c in cols}
    for r in raw[1:]:
        row = {c.key: "" for c in cols}
        for i, key in mapping.items():
            if i < len(r):
                v = r[i]
                if isinstance(v, datetime):
                    v = v.date().isoformat()
                elif isinstance(v, date):
                    v = v.isoformat()
                elif isinstance(v, float) and v.is_integer():
                    v = str(int(v))
                row[key] = str(v).strip()
        if all(row[k] == example[k] for k in row):
            continue  # the untouched example row from the template
        rows.append(row)
    if len(rows) > MAX_ROWS:
        raise ImportFileError(f"Up to {MAX_ROWS} rows per import. Split the file and import it in parts.")
    if not rows:
        raise ImportFileError("No data rows found under the column names")
    return rows


# ------------------------------------------------------------------ validation
@dataclass
class RowResult:
    row: int
    data: dict
    action: str = "create"   # create | update
    errors: list = field(default_factory=list)
    notes: list = field(default_factory=list)

    def err(self, f: str, msg: str):
        self.errors.append({"field": f, "message": msg})


def _dec(r: RowResult, f: str, *, minimum=Decimal(0), positive=False) -> Decimal | None:
    v = (r.data.get(f) or "").replace(",", "").strip()
    if v == "":
        return None
    try:
        d = Decimal(v)
    except InvalidOperation:
        r.err(f, "Must be a number")
        return None
    if positive and d <= 0:
        r.err(f, "Must be more than 0")
    elif d < minimum:
        r.err(f, "Can't be negative")
    return d


def _bool(r: RowResult, f: str, default=None):
    v = (r.data.get(f) or "").strip().lower()
    if v == "":
        return default
    if v in ("yes", "y", "true", "1", "x"):
        return True
    if v in ("no", "n", "false", "0"):
        return False
    r.err(f, "Use yes or no")
    return default


def _date(r: RowResult, f: str) -> date | None:
    v = (r.data.get(f) or "").strip()
    if not v:
        return None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y"):
        try:
            return datetime.strptime(v, fmt).date()
        except ValueError:
            continue
    r.err(f, "Use a date like 2027-03-31")
    return None


def _clean(rows: list[dict], kind: str) -> list[dict]:
    keys = [c.key for c in KINDS[kind]["columns"]]
    return [{k: str(row.get(k) or "").strip() for k in keys} for row in rows]


def validate(db: Session, ctx: inv.Ctx, kind: str, rows: list[dict]) -> list[RowResult]:
    rows = _clean(rows, kind)
    if len(rows) > MAX_ROWS:
        raise ImportFileError(f"Up to {MAX_ROWS} rows per import")
    results = [RowResult(row=i + 2, data=r) for i, r in enumerate(rows)]  # +2: header is row 1
    for r in results:
        for c in KINDS[kind]["columns"]:
            if c.required and not r.data.get(c.key):
                r.err(c.key, f"{c.label} is required")
    {"products": _v_products, "suppliers": _v_partners, "customers": _v_partners,
     "opening-stock": _v_opening}[kind](db, ctx, kind, results)
    return results


def _v_products(db, ctx, kind, results):
    org = ctx.org_id
    units = {}
    for u in db.execute(select(UnitOfMeasure).where(UnitOfMeasure.org_id == org)).scalars():
        units[_norm(u.abbreviation)] = u
        units[_norm(u.name)] = u
    suppliers = {s.name.strip().lower(): s for s in db.execute(select(Supplier).where(Supplier.org_id == org)).scalars()}
    skus = [r.data["sku"] for r in results if r.data["sku"]]
    existing = {p.sku.lower(): p for p in db.execute(select(Product).where(
        Product.org_id == org, func.lower(Product.sku).in_([s.lower() for s in skus]))).unique().scalars()}
    barcodes = [r.data["barcode"] for r in results if r.data["barcode"]]
    taken = {b: s for b, s in db.execute(select(Product.barcode, Product.sku).where(
        Product.org_id == org, Product.barcode.in_(barcodes))).all()} if barcodes else {}
    seen_sku, seen_bc = {}, {}
    for r in results:
        d = r.data
        sku = d["sku"]
        if sku:
            if sku.lower() in seen_sku:
                r.err("sku", f"Same SKU as row {seen_sku[sku.lower()]}")
            seen_sku[sku.lower()] = r.row
            if len(sku) > 64:
                r.err("sku", "Up to 64 characters")
        p = existing.get(sku.lower())
        r.action = "update" if p else "create"
        if p and p.has_variants:
            r.err("sku", "This SKU is a product with sizes/colours; import the individual sizes/colours instead")
        if d["barcode"]:
            if d["barcode"] in seen_bc:
                r.err("barcode", f"Same barcode as row {seen_bc[d['barcode']]}")
            seen_bc[d["barcode"]] = r.row
            owner = taken.get(d["barcode"])
            if owner and owner.lower() != sku.lower():
                r.err("barcode", f"Already used by product {owner}")
        if d["unit"]:
            u = units.get(_norm(d["unit"]))
            if not u:
                r.err("unit", f"Unknown unit '{d['unit']}'. Use one of: {', '.join(sorted({x.abbreviation for x in units.values()}))}")
            elif p and u.id != p.base_unit_id:
                r.err("unit", "The unit of an existing product can't be changed by import")
        if d["supplier"] and d["supplier"].lower() not in suppliers:
            r.err("supplier", f"No supplier called '{d['supplier']}'. Import suppliers first or leave this empty")
        for f in ("cost_price", "selling_price", "reorder_level", "reorder_qty", "min_stock"):
            _dec(r, f)
        tax = _dec(r, "tax_rate")
        if tax is not None and tax > 100:
            r.err("tax_rate", "Tax can't be more than 100%")
        flags = {f: _bool(r, f) for f in ("track_expiry", "track_batches", "track_serials", "is_active")}
        if p:
            changing = [f for f in ("track_batches", "track_serials") if flags[f] is not None and flags[f] != getattr(p, f)]
            if flags["track_expiry"] and not p.track_batches:
                changing.append("track_expiry")
            if changing and db.execute(select(StockMovement.id).where(StockMovement.product_id == p.id).limit(1)).first():
                r.err(changing[0], "Tracking can't change on a product that already has stock history")
        if d["category"] and not p:
            if not db.execute(select(Category.id).where(Category.org_id == org, func.lower(Category.name) == d["category"].lower())).first():
                r.notes.append(f"New category '{d['category']}' will be created")


def _v_partners(db, ctx, kind, results):
    model = Supplier if kind == "suppliers" else Customer
    existing = {x.name.strip().lower() for x in db.execute(select(model).where(model.org_id == ctx.org_id)).scalars()}
    seen = {}
    for r in results:
        name = r.data["name"]
        if name:
            k = name.lower()
            if k in seen:
                r.err("name", f"Same name as row {seen[k]}")
            seen[k] = r.row
            r.action = "update" if k in existing else "create"
            if len(name) > 200:
                r.err("name", "Up to 200 characters")
        if r.data["email"] and not EMAIL_RE.match(r.data["email"]):
            r.err("email", "Not a valid email address")


def _v_opening(db, ctx, kind, results):
    org = ctx.org_id
    skus = {r.data["sku"].lower() for r in results if r.data["sku"]}
    products = {p.sku.lower(): p for p in db.execute(select(Product).where(
        Product.org_id == org, func.lower(Product.sku).in_(skus))).unique().scalars()} if skus else {}
    whs = db.execute(select(Warehouse).where(Warehouse.org_id == org, Warehouse.is_active.is_(True))).scalars().all()
    by_wh = {}
    for w in whs:
        by_wh[w.name.strip().lower()] = w
        by_wh[w.code.lower()] = w
    locs = db.execute(select(Location).where(Location.org_id == org, Location.is_active.is_(True))).unique().scalars().all()
    serial_seen = {}
    for r in results:
        d = r.data
        p = products.get(d["sku"].lower()) if d["sku"] else None
        if d["sku"] and not p:
            r.err("sku", f"No product with SKU '{d['sku']}'. Import products first")
        elif p and p.has_variants:
            r.err("sku", f"{p.name} has sizes/colours; use the SKU of the exact size/colour")
        w = by_wh.get(d["location"].lower()) if d["location"] else None
        if d["location"] and not w:
            r.err("location", f"No location called '{d['location']}'")
        loc = None
        if w:
            if ctx.allowed_warehouse_ids and w.id not in ctx.allowed_warehouse_ids:
                r.err("location", "You are not allowed to work in this location")
            if d["bin"]:
                loc = next((l for l in locs if l.warehouse_id == w.id and l.code.lower() == d["bin"].lower()), None)
                if not loc:
                    r.err("bin", f"No shelf/bin '{d['bin']}' in {w.name}")
            else:
                loc = next((l for l in locs if l.warehouse_id == w.id and l.kind == LocationKind.DEFAULT), None)
        r.data["_location_id"] = str(loc.id) if loc else ""
        qty = _dec(r, "quantity", positive=True)
        _dec(r, "unit_cost")
        exp = _date(r, "expires_on")
        if p and qty is not None and qty > 0:
            if not p.base_unit.allow_decimals and qty != qty.to_integral_value():
                r.err("quantity", f"{p.name} is counted in whole {p.base_unit.name.lower()}s")
            if p.track_batches and not d["batch_number"]:
                r.err("batch_number", f"{p.name} needs a batch number")
            if p.track_expiry and not exp and not d["expires_on"]:
                r.err("expires_on", f"{p.name} needs an expiry date")
            if p.track_serials:
                serials = [s.strip() for s in re.split(r"[;,\n]", d["serials"]) if s.strip()]
                if Decimal(len(serials)) != qty:
                    r.err("serials", f"List {inv.fmt(qty)} serial number(s), separated by ;")
                for s in serials:
                    if (p.id, s) in serial_seen:
                        r.err("serials", f"Serial {s} is also on row {serial_seen[(p.id, s)]}")
                    serial_seen[(p.id, s)] = r.row
        r.action = "create"


# ------------------------------------------------------------------ commit
def commit(db: Session, ctx: inv.Ctx, kind: str, rows: list[dict]) -> dict:
    results = validate(db, ctx, kind, rows)
    if any(r.errors for r in results):
        return {"ok": False, "results": results}
    counts = {"created": 0, "updated": 0}
    if kind == "products":
        _c_products(db, ctx, results, counts)
    elif kind in ("suppliers", "customers"):
        _c_partners(db, ctx, Supplier if kind == "suppliers" else Customer, results, counts)
    else:
        _c_opening(db, ctx, results, counts)
    db.flush()
    return {"ok": True, "results": results, **counts}


def _c_products(db, ctx, results, counts):
    org = ctx.org_id
    units = {}
    for u in db.execute(select(UnitOfMeasure).where(UnitOfMeasure.org_id == org)).scalars():
        units[_norm(u.abbreviation)] = u
        units[_norm(u.name)] = u
    pcs = units.get("pcs")
    suppliers = {s.name.strip().lower(): s for s in db.execute(select(Supplier).where(Supplier.org_id == org)).scalars()}
    cats = {c.name.lower(): c for c in db.execute(select(Category).where(Category.org_id == org)).scalars()}
    defaults = product_defaults(db.get(Organization, org).settings or {})
    for r in results:
        d = r.data
        p = db.execute(select(Product).where(Product.org_id == org, func.lower(Product.sku) == d["sku"].lower())).unique().scalar_one_or_none()
        if p is None:
            p = Product(org_id=org, sku=d["sku"], name=d["name"], base_unit_id=(units.get(_norm(d["unit"])) or pcs).id,
                        track_batches=defaults["track_batches"], track_expiry=defaults["track_expiry"], track_serials=False)
            db.add(p)
            counts["created"] += 1
        else:
            counts["updated"] += 1
        p.name = d["name"] or p.name
        for f in ("barcode", "brand", "description"):
            if d[f]:
                setattr(p, f, d[f])
        if d["category"]:
            c = cats.get(d["category"].lower())
            if not c:
                c = Category(org_id=org, name=d["category"])
                db.add(c)
                db.flush()
                cats[c.name.lower()] = c
            p.category_id = c.id
        if d["supplier"]:
            p.preferred_supplier_id = suppliers[d["supplier"].lower()].id
        for f in ("cost_price", "selling_price", "tax_rate", "reorder_level", "reorder_qty", "min_stock"):
            if d[f]:
                setattr(p, f, Decimal(d[f].replace(",", "")))
        if r.action == "create":
            p.average_cost = p.cost_price or 0
        for f in ("track_batches", "track_serials", "is_active"):
            v = _bool(r, f)
            if v is not None:
                setattr(p, f, v)
        te = _bool(r, "track_expiry")
        if te is not None:
            p.track_expiry = te
            if te:
                p.track_batches = True
        db.flush()


def _c_partners(db, ctx, model, results, counts):
    existing = {x.name.strip().lower(): x for x in db.execute(select(model).where(model.org_id == ctx.org_id)).scalars()}
    for r in results:
        d = r.data
        obj = existing.get(d["name"].lower())
        if obj is None:
            obj = model(org_id=ctx.org_id, name=d["name"])
            db.add(obj)
            existing[d["name"].lower()] = obj
            counts["created"] += 1
        else:
            counts["updated"] += 1
        for f in ("contact_name", "email", "phone", "address", "notes"):
            if d[f]:
                setattr(obj, f, d[f])


def _c_opening(db, ctx, results, counts):
    for r in results:
        d = r.data
        p = db.execute(select(Product).where(Product.org_id == ctx.org_id, func.lower(Product.sku) == d["sku"].lower())).unique().scalar_one()
        serials = [s.strip() for s in re.split(r"[;,\n]", d["serials"]) if s.strip()] if p.track_serials else None
        inv.receive(db, ctx, product_id=p.id, location_id=d["_location_id"], quantity=Decimal(d["quantity"].replace(",", "")),
                    unit_cost=Decimal(d["unit_cost"].replace(",", "")) if d["unit_cost"] else p.cost_price,
                    batch_number=d["batch_number"] or None, expires_on=_date(RowResult(0, d), "expires_on"),
                    serials=serials, mtype=MovementType.OPENING_BALANCE, note="Opening stock (import)",
                    reference_type="import")
        counts["created"] += 1
