"""Turn a tabular report into CSV, Excel or PDF."""
import csv
import io
from datetime import date, datetime
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle

CONTENT_TYPES = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
}


def _cell(v):
    if v is None:
        return ""
    if isinstance(v, Decimal):
        return float(v) if v != v.to_integral_value() else int(v)
    if isinstance(v, (date, datetime)):
        return v.strftime("%Y-%m-%d %H:%M") if isinstance(v, datetime) else v.isoformat()
    if hasattr(v, "value"):
        return v.value
    return v


def _safe_csv(v):
    # Prevent spreadsheet formula injection when a CSV is opened in Excel.
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + v
    return v


def to_csv(columns: list[dict], rows: list[dict]) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([c["label"] for c in columns])
    for r in rows:
        w.writerow([_safe_csv(_cell(r.get(c["key"]))) for c in columns])
    return ("﻿" + buf.getvalue()).encode("utf-8")  # BOM so Excel reads UTF-8


def to_xlsx(title: str, columns: list[dict], rows: list[dict], subtitle: str = "") -> bytes:
    wb = Workbook(write_only=False)
    ws = wb.active
    ws.title = title[:31]
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([subtitle])
    ws.append([])
    ws.append([c["label"] for c in columns])
    head_fill = PatternFill("solid", fgColor="E8EEF6")
    for i in range(1, len(columns) + 1):
        cell = ws.cell(row=4, column=i)
        cell.font = Font(bold=True)
        cell.fill = head_fill
        cell.alignment = Alignment(vertical="center")
    for r in rows:
        vals = []
        for c in columns:
            v = _cell(r.get(c["key"]))
            if isinstance(v, str) and v[:1] == "=":
                v = "'" + v
            vals.append(v)
        ws.append(vals)
    for i, c in enumerate(columns, 1):
        width = max([len(str(c["label"]))] + [len(str(_cell(r.get(c["key"])))) for r in rows[:500]]) + 2
        ws.column_dimensions[get_column_letter(i)].width = min(max(width, 8), 50)
        if c.get("type") in ("money", "number"):
            for row in ws.iter_rows(min_row=5, min_col=i, max_col=i):
                row[0].number_format = "#,##0.00" if c["type"] == "money" else "#,##0.####"
    ws.freeze_panes = "A5"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def to_pdf(title: str, columns: list[dict], rows: list[dict], subtitle: str = "") -> bytes:
    out = io.BytesIO()
    doc = SimpleDocTemplate(out, pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm,
                            topMargin=12 * mm, bottomMargin=12 * mm, title=title)
    styles = getSampleStyleSheet()
    small = styles["BodyText"].clone("small", fontSize=7.5, leading=9)
    head = styles["BodyText"].clone("head", fontSize=7.5, leading=9, fontName="Helvetica-Bold")
    data = [[Paragraph(c["label"], head) for c in columns]]
    for r in rows:
        data.append([Paragraph(str(_cell(r.get(c["key"]))).replace("&", "&amp;").replace("<", "&lt;"), small) for c in columns])
    table = LongTable(data, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EEF6")),
        ("LINEBELOW", (0, 0), (-1, 0), 0.6, colors.HexColor("#94A3B8")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    story = [Paragraph(title, styles["Title"]), Paragraph(subtitle, styles["Normal"]), Spacer(1, 6 * mm), table]
    if not rows:
        story.append(Paragraph("No rows.", styles["Normal"]))
    doc.build(story)
    return out.getvalue()


def render(fmt: str, title: str, columns: list[dict], rows: list[dict], subtitle: str = "") -> bytes:
    if fmt == "csv":
        return to_csv(columns, rows)
    if fmt == "xlsx":
        return to_xlsx(title, columns, rows, subtitle)
    if fmt == "pdf":
        return to_pdf(title, columns, rows, subtitle)
    raise ValueError(fmt)
